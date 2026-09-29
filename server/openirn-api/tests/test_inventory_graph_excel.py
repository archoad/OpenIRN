from __future__ import annotations

import sqlite3
import sys
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

import main as api  # noqa: E402


ASSET_A = "11111111-1111-4111-8111-111111111111"
SYSTEM_A = "22222222-2222-4222-8222-222222222222"
FUNCTION_A = "33333333-3333-4333-8333-333333333333"
FOREIGN_ASSET = "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa"


def _database() -> sqlite3.Connection:
    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.executescript(
        """
        CREATE TABLE tenants(
            id TEXT NOT NULL PRIMARY KEY,
            display_name TEXT NOT NULL
        );
        CREATE TABLE critical_functions(
            tenant_id TEXT NOT NULL,
            function_id TEXT NOT NULL,
            name TEXT NOT NULL,
            description TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            PRIMARY KEY(tenant_id, function_id)
        );
        CREATE TABLE information_systems(
            tenant_id TEXT NOT NULL,
            system_id TEXT NOT NULL,
            function_id TEXT,
            name TEXT NOT NULL,
            description TEXT NOT NULL,
            owner TEXT NOT NULL,
            owner_first_name TEXT NOT NULL,
            owner_last_name TEXT NOT NULL,
            owner_email TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            PRIMARY KEY(tenant_id, system_id)
        );
        CREATE TABLE information_assets(
            tenant_id TEXT NOT NULL,
            asset_id TEXT NOT NULL,
            system_id TEXT,
            name TEXT NOT NULL,
            asset_type TEXT NOT NULL,
            description TEXT NOT NULL,
            criticality TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            PRIMARY KEY(tenant_id, asset_id)
        );
        CREATE TABLE critical_function_systems(
            tenant_id TEXT NOT NULL,
            function_id TEXT NOT NULL,
            system_id TEXT NOT NULL,
            created_at TEXT NOT NULL,
            PRIMARY KEY(tenant_id, function_id, system_id)
        );
        CREATE TABLE information_system_assets(
            tenant_id TEXT NOT NULL,
            system_id TEXT NOT NULL,
            asset_id TEXT NOT NULL,
            created_at TEXT NOT NULL,
            PRIMARY KEY(tenant_id, system_id, asset_id)
        );
        CREATE TABLE asset_assessment_answers(
            tenant_id TEXT NOT NULL,
            asset_id TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        """
    )
    now = "2026-09-24T12:00:00+00:00"
    con.execute("INSERT INTO tenants VALUES (?, ?)", ("tenant-a", "Espace A"))
    con.execute("INSERT INTO tenants VALUES (?, ?)", ("tenant-b", "Espace B"))
    con.execute(
        "INSERT INTO critical_functions VALUES (?, ?, ?, ?, ?, ?)",
        ("tenant-a", FUNCTION_A, "Continuité métier", "Fonction cœur", now, now),
    )
    con.execute(
        "INSERT INTO information_systems VALUES (?, ?, NULL, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            "tenant-a",
            SYSTEM_A,
            "SI Échanges",
            "Description Unicode — 日本語",
            "Élodie Martin",
            "Élodie",
            "Martin",
            "elodie@example.test",
            now,
            now,
        ),
    )
    con.execute(
        "INSERT INTO information_assets VALUES (?, ?, NULL, ?, ?, ?, ?, ?, ?)",
        (
            "tenant-a",
            ASSET_A,
            "Base clients",
            "Base de données",
            "Données à caractère personnel",
            "4",
            now,
            now,
        ),
    )
    con.execute(
        "INSERT INTO information_assets VALUES (?, ?, NULL, ?, ?, ?, ?, ?, ?)",
        (
            "tenant-b",
            FOREIGN_ASSET,
            "Actif étranger",
            "Service",
            "Autre espace",
            "2",
            now,
            now,
        ),
    )
    con.execute(
        "INSERT INTO critical_function_systems VALUES (?, ?, ?, ?)",
        ("tenant-a", FUNCTION_A, SYSTEM_A, now),
    )
    con.execute(
        "INSERT INTO information_system_assets VALUES (?, ?, ?, ?)",
        ("tenant-a", SYSTEM_A, ASSET_A, now),
    )
    con.commit()
    return con


def _edit_workbook(raw: bytes, editor) -> bytes:
    workbook = load_workbook(BytesIO(raw))
    try:
        editor(workbook)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


class InventoryGraphExcelTests(unittest.TestCase):
    def setUp(self) -> None:
        self.con = _database()
        self.raw = api._inventory_graph_to_excel_bytes(self.con, "tenant-a")

    def tearDown(self) -> None:
        self.con.close()

    def test_round_trip_preserves_unicode_catalogues_and_links(self) -> None:
        counts = api._inventory_graph_import_from_excel_bytes(
            self.con,
            "tenant-a",
            self.raw,
        )

        self.assertEqual(counts["assets"], 1)
        self.assertEqual(counts["informationSystems"], 1)
        self.assertEqual(counts["criticalFunctions"], 1)
        self.assertEqual(counts["assetSystemLinks"], 1)
        self.assertEqual(counts["systemFunctionLinks"], 1)
        system = self.con.execute(
            "SELECT name, description, owner_first_name FROM information_systems WHERE tenant_id = ?",
            ("tenant-a",),
        ).fetchone()
        self.assertEqual(system["name"], "SI Échanges")
        self.assertEqual(system["description"], "Description Unicode — 日本語")
        self.assertEqual(system["owner_first_name"], "Élodie")

    def test_default_key_is_a_readable_slug_not_the_raw_uuid(self) -> None:
        workbook = load_workbook(BytesIO(self.raw))
        try:
            self.assertEqual(workbook["Actifs"]["A2"].value, "base-clients")
            self.assertEqual(workbook["Actifs"]["F2"].value, ASSET_A)
            self.assertEqual(workbook["SI"]["A2"].value, "si-echanges")
            self.assertEqual(workbook["SI"]["G2"].value, SYSTEM_A)
            self.assertEqual(
                workbook["Fonctions critiques"]["A2"].value, "continuite-metier"
            )
            self.assertEqual(workbook["Fonctions critiques"]["D2"].value, FUNCTION_A)
            self.assertEqual(
                workbook["Actifs - SI"]["A2"].value, "base-clients"
            )
            self.assertEqual(workbook["Actifs - SI"]["B2"].value, "si-echanges")
        finally:
            workbook.close()

    def test_duplicate_names_get_deduplicated_keys(self) -> None:
        now = "2026-09-24T12:00:00+00:00"
        self.con.execute(
            "INSERT INTO information_assets VALUES (?, ?, NULL, ?, ?, ?, ?, ?, ?)",
            ("tenant-a", "44444444-4444-4444-8444-444444444444", "Base clients", "Base de données", "", "2", now, now),
        )
        self.con.commit()
        raw = api._inventory_graph_to_excel_bytes(self.con, "tenant-a")
        workbook = load_workbook(BytesIO(raw))
        try:
            keys = {workbook["Actifs"]["A2"].value, workbook["Actifs"]["A3"].value}
            self.assertEqual(keys, {"base-clients", "base-clients-2"})
        finally:
            workbook.close()

    def test_long_generated_keys_are_bounded_deduplicated_and_round_trip(self) -> None:
        now = "2026-09-24T12:00:00+00:00"
        second_asset_id = "44444444-4444-4444-8444-444444444444"
        self.con.execute(
            "UPDATE information_assets SET name = ? WHERE tenant_id = ? AND asset_id = ?",
            (f"{'a' * 200}x", "tenant-a", ASSET_A),
        )
        self.con.execute(
            "UPDATE information_systems SET name = ? WHERE tenant_id = ? AND system_id = ?",
            ("s" * 200, "tenant-a", SYSTEM_A),
        )
        self.con.execute(
            "UPDATE critical_functions SET name = ? WHERE tenant_id = ? AND function_id = ?",
            ("f" * 200, "tenant-a", FUNCTION_A),
        )
        self.con.execute(
            "INSERT INTO information_assets VALUES (?, ?, NULL, ?, ?, ?, ?, ?, ?)",
            (
                "tenant-a",
                second_asset_id,
                f"{'a' * 200}y",
                "Service",
                "",
                "2",
                now,
                now,
            ),
        )
        self.con.execute(
            "INSERT INTO information_system_assets VALUES (?, ?, ?, ?)",
            ("tenant-a", SYSTEM_A, second_asset_id, now),
        )
        self.con.commit()

        raw = api._inventory_graph_to_excel_bytes(self.con, "tenant-a")
        workbook = load_workbook(BytesIO(raw))
        try:
            asset_keys = [workbook["Actifs"][f"A{row}"].value for row in (2, 3)]
            self.assertEqual(len(set(asset_keys)), 2)
            self.assertTrue(all(len(key) <= 160 for key in asset_keys))
            self.assertTrue(any(key.endswith("-2") for key in asset_keys))
            self.assertLessEqual(len(workbook["SI"]["A2"].value), 160)
            self.assertLessEqual(
                len(workbook["Fonctions critiques"]["A2"].value),
                160,
            )
        finally:
            workbook.close()

        counts = api._inventory_graph_import_from_excel_bytes(
            self.con,
            "tenant-a",
            raw,
        )
        self.assertEqual(counts["assets"], 2)
        self.assertEqual(counts["assetSystemLinks"], 2)

    def test_custom_key_longer_than_limit_is_rejected_before_write(self) -> None:
        oversized_key = "k" * 161

        def replace_key(workbook) -> None:
            workbook["Actifs"]["A2"] = oversized_key
            workbook["Actifs - SI"]["A2"] = oversized_key

        raw = _edit_workbook(self.raw, replace_key)
        with self.assertRaisesRegex(HTTPException, "limitée à 160 caractères"):
            api._inventory_graph_import_from_excel_bytes(
                self.con,
                "tenant-a",
                raw,
            )

        stored_name = self.con.execute(
            "SELECT name FROM information_assets WHERE tenant_id = ? AND asset_id = ?",
            ("tenant-a", ASSET_A),
        ).fetchone()[0]
        self.assertEqual(stored_name, "Base clients")

    def test_formula_like_inventory_values_are_exported_as_literal_text(self) -> None:
        formula_like_name = "=HYPERLINK(\"https://example.test\",\"Actif\")"
        self.con.execute(
            "UPDATE information_assets SET name = ? WHERE tenant_id = ? AND asset_id = ?",
            (formula_like_name, "tenant-a", ASSET_A),
        )
        self.con.commit()

        raw = api._inventory_graph_to_excel_bytes(self.con, "tenant-a")
        workbook = load_workbook(BytesIO(raw), data_only=False)
        try:
            for sheet_name, coordinate in (
                ("Actifs", "B2"),
                ("Aperçu actifs", "A2"),
            ):
                cell = workbook[sheet_name][coordinate]
                self.assertEqual(cell.value, formula_like_name)
                self.assertEqual(cell.data_type, "s")
        finally:
            workbook.close()

        scoped_raw = api._inventory_to_excel_bytes(
            self.con,
            "tenant-a",
            SYSTEM_A,
        )
        scoped_workbook = load_workbook(BytesIO(scoped_raw), data_only=False)
        try:
            cell = scoped_workbook["Actifs SI"]["B2"]
            self.assertEqual(cell.value, formula_like_name)
            self.assertEqual(cell.data_type, "s")
        finally:
            scoped_workbook.close()

        api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)
        stored_name = self.con.execute(
            "SELECT name FROM information_assets WHERE tenant_id = ? AND asset_id = ?",
            ("tenant-a", ASSET_A),
        ).fetchone()[0]
        self.assertEqual(stored_name, formula_like_name)

    def test_overview_sheets_are_read_only_and_human_readable(self) -> None:
        workbook = load_workbook(BytesIO(self.raw))
        try:
            asset_overview = workbook["Aperçu actifs"]
            self.assertEqual(
                [cell.value for cell in asset_overview[1]],
                ["Nom actif", "Type actif", "Criticité actif", "SI associés"],
            )
            self.assertEqual(
                [cell.value for cell in asset_overview[2]],
                ["Base clients", "Base de données", "4", "SI Échanges"],
            )
            system_overview = workbook["Aperçu SI"]
            self.assertEqual(
                [cell.value for cell in system_overview[1]],
                ["Nom SI", "Directeur", "Fonctions critiques associées"],
            )
            self.assertEqual(
                [cell.value for cell in system_overview[2]],
                ["SI Échanges", "Élodie Martin", "Continuité métier"],
            )
        finally:
            workbook.close()

    def test_overview_sheets_are_ignored_on_import(self) -> None:
        counts = api._inventory_graph_import_from_excel_bytes(
            self.con,
            "tenant-a",
            self.raw,
        )
        self.assertEqual(counts["assets"], 1)

    def test_blank_template_rows_are_ignored(self) -> None:
        counts = api._inventory_graph_import_from_excel_bytes(
            self.con,
            "tenant-a",
            self.raw,
        )
        self.assertEqual(counts["assets"], 1)

    def test_missing_required_column_is_rejected_before_write(self) -> None:
        raw = _edit_workbook(
            self.raw,
            lambda workbook: setattr(workbook["Actifs"]["A1"], "value", ""),
        )
        with self.assertRaisesRegex(HTTPException, "colonne obligatoire absente"):
            api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)
        self.assertEqual(
            self.con.execute(
                "SELECT COUNT(*) FROM information_system_assets WHERE tenant_id = ?",
                ("tenant-a",),
            ).fetchone()[0],
            1,
        )

    def test_unknown_header_is_rejected(self) -> None:
        raw = _edit_workbook(
            self.raw,
            lambda workbook: setattr(workbook["Actifs"]["G1"], "value", "Colonne surprise"),
        )
        with self.assertRaisesRegex(HTTPException, "en-tête inconnu"):
            api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)

    def test_duplicate_identifier_is_rejected(self) -> None:
        def duplicate(workbook) -> None:
            workbook["Actifs"].append(
                ["autre-cle", "Copie", "Service", "2", "Dupliqué", ASSET_A]
            )

        raw = _edit_workbook(self.raw, duplicate)
        with self.assertRaisesRegex(HTTPException, "ID actif dupliqué"):
            api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)

    def test_identifier_from_another_tenant_is_rejected(self) -> None:
        raw = _edit_workbook(
            self.raw,
            lambda workbook: setattr(workbook["Actifs"]["F2"], "value", FOREIGN_ASSET),
        )
        with self.assertRaisesRegex(HTTPException, "identifiant inconnu dans cet espace"):
            api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)

    def test_invalid_criticality_is_rejected(self) -> None:
        raw = _edit_workbook(
            self.raw,
            lambda workbook: setattr(workbook["Actifs"]["D2"], "value", "5"),
        )
        with self.assertRaisesRegex(HTTPException, "comprise entre 1 et 4"):
            api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)

    def test_new_rows_with_business_keys_create_and_link_entities(self) -> None:
        def add_asset(workbook) -> None:
            system_key = workbook["SI"]["A2"].value
            workbook["Actifs"].append(
                ["nouvel-actif", "API partenaires", "Service", "3", "Créé par import", ""]
            )
            workbook["Actifs - SI"].append(["nouvel-actif", system_key])

        raw = _edit_workbook(self.raw, add_asset)
        counts = api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)
        self.assertEqual(counts["createdAssets"], 1)
        self.assertEqual(counts["assetSystemLinks"], 2)
        self.assertEqual(
            self.con.execute(
                "SELECT COUNT(*) FROM information_assets WHERE tenant_id = ?",
                ("tenant-a",),
            ).fetchone()[0],
            2,
        )

    def test_blank_keys_create_unlinked_items_and_existing_labels_are_updated(self) -> None:
        def update_and_add_items(workbook) -> None:
            workbook["Actifs"]["B2"] = "Base clients renommée"
            workbook["SI"]["B2"] = "SI Échanges renommé"
            workbook["Fonctions critiques"]["B2"] = "Continuité renommée"
            workbook["Actifs"].append(
                ["", "Cluster ELK", "Serveurs", "1", "Créé sans clé", ""]
            )
            workbook["SI"].append(
                ["", "SI Observabilité", "Nouveau SI", "Léa", "Durand", "lea@example.test", ""]
            )
            workbook["Fonctions critiques"].append(
                ["", "Supervision", "Nouvelle fonction", ""]
            )

        raw = _edit_workbook(self.raw, update_and_add_items)
        counts = api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)

        self.assertEqual(counts["createdAssets"], 1)
        self.assertEqual(counts["createdInformationSystems"], 1)
        self.assertEqual(counts["createdCriticalFunctions"], 1)
        self.assertEqual(
            self.con.execute(
                "SELECT name FROM information_assets WHERE tenant_id = ? AND asset_id = ?",
                ("tenant-a", ASSET_A),
            ).fetchone()[0],
            "Base clients renommée",
        )
        self.assertEqual(
            self.con.execute(
                "SELECT name FROM information_systems WHERE tenant_id = ? AND system_id = ?",
                ("tenant-a", SYSTEM_A),
            ).fetchone()[0],
            "SI Échanges renommé",
        )
        self.assertEqual(
            self.con.execute(
                "SELECT name FROM critical_functions WHERE tenant_id = ? AND function_id = ?",
                ("tenant-a", FUNCTION_A),
            ).fetchone()[0],
            "Continuité renommée",
        )

    def test_export_filename_uses_requested_date_prefix(self) -> None:
        fixed_now = datetime(2026, 9, 25, 8, 30, tzinfo=timezone.utc)
        with (
            patch.object(api, "_require_campaign_manager_authorization"),
            patch.object(api, "_db", return_value=self.con),
            patch.object(api, "_ensure_tenant"),
            patch.object(api, "_utc_now", return_value=fixed_now),
        ):
            response = api.asset_inventory_graph_export_excel(
                object(),
                tenantId="tenant-a",
            )

        self.assertEqual(
            response.headers["content-disposition"],
            'attachment; filename="20260925_openirn_export.xlsx"',
        )

    def test_incomplete_catalog_is_rejected(self) -> None:
        raw = _edit_workbook(
            self.raw,
            lambda workbook: workbook["Actifs"].delete_rows(2),
        )
        with self.assertRaisesRegex(HTTPException, "doit conserver tous les éléments"):
            api._inventory_graph_import_from_excel_bytes(self.con, "tenant-a", raw)

    def test_large_workbook_is_rejected(self) -> None:
        original = api._excel_rows_strict

        def oversized(sheet, **kwargs):
            rows = original(sheet, **kwargs)
            if sheet.title == "Actifs":
                return rows * 20001
            return rows

        with patch.object(api, "_excel_rows_strict", side_effect=oversized):
            with self.assertRaisesRegex(HTTPException, "trop d'éléments"):
                api._inventory_graph_import_from_excel_bytes(
                    self.con,
                    "tenant-a",
                    self.raw,
                )


if __name__ == "__main__":
    unittest.main()
